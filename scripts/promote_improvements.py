"""Move the cloud baseline forward for improvements only.

After a run is compared with the baseline, each component/locale that improved is copied
from the run into the baseline, so the improvement stops being reported as a change.
Everything else stays at the old baseline, so a regression keeps being reported every run
until it is fixed.

A component/locale is promoted only when ALL of these hold (from DifferenceReport.xlsx):
  - Direction is "Improved" (overall match % went up)
  - no change made the comparison worse (no "(worse)" row in Details)
  - no page error, and the status did not drop (e.g. WARN -> FAIL)

Writes daily/<run>/BaselineUpdate.txt and appends a line to BASELINE_SOURCE.txt.

Usage: python3 scripts/promote_improvements.py <run_dir> [--baseline baselines/cloud_baseline] [--dry-run]
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

import generate_html_report

ROOT = Path(__file__).resolve().parent.parent
CLOUD_BASELINE_DIR = ROOT / "baselines" / "cloud_baseline"
COMPONENTS = [("Header", "HeaderCheck_Report"), ("Footer", "FooterCheck_Report")]
STATUS_RANK = {"ERROR": 0, "FAIL": 1, "WARN": 2, "PASS": 3}


def _rows(path: Path, sheet: str) -> list[dict]:
    wb = load_workbook(path, read_only=True)
    if sheet not in wb.sheetnames:
        return []
    rows = list(wb[sheet].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:] if any(v is not None for v in r)]


def decide(run_dir: Path) -> tuple[list[tuple[str, str]], list[tuple[str, str, str]]]:
    """Return (promote, held_back). held_back items carry the reason."""
    diff = run_dir / "DifferenceReport.xlsx"
    if not diff.exists():
        return [], []
    worse = {(r["Component"], r["Locale"]) for r in _rows(diff, "Details")
             if "(worse)" in str(r["Change"]) or r["Page"] == "Run"}
    promote, held = [], []
    for r in _rows(diff, "Overview"):
        key = (r["Component"], r["Locale"])
        if r["Direction"] != "Improved":
            continue
        if key in worse:
            held.append((*key, "also has a change that made it worse"))
        elif STATUS_RANK.get(r["Current Status"], 0) < STATUS_RANK.get(r["Baseline Status"], 0):
            held.append((*key, f"status dropped {r['Baseline Status']} -> {r['Current Status']}"))
        else:
            promote.append(key)
    return promote, held


def merge_workbook(base_path: Path, cur_path: Path, locales: set[str]) -> None:
    """Replace the rows of `locales` in every per-locale sheet of base_path with those of cur_path."""
    base = load_workbook(base_path)
    cur = load_workbook(cur_path, read_only=True)
    for ws in base.worksheets:
        headers = [c.value for c in ws[1]]
        if "Locale" not in headers or ws.title not in cur.sheetnames:
            continue
        li = headers.index("Locale")
        cur_rows = [r for r in list(cur[ws.title].iter_rows(values_only=True))[1:] if r and r[li] in locales]
        new, done = [], set()
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not any(v is not None for v in r):
                continue
            loc = r[li]
            if loc in locales:
                if loc not in done:  # keep the locale's position; insert the run's rows once
                    new += [c for c in cur_rows if c[li] == loc]
                    done.add(loc)
            else:
                new.append(r)
        new += [c for c in cur_rows if c[li] not in done]  # locale had no rows in this sheet before
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for r in new:
            ws.append(list(r))
    base.save(base_path)


def run(run_dir: Path, baseline_dir: Path = CLOUD_BASELINE_DIR, dry_run: bool = False) -> list[tuple[str, str]]:
    promote, held = decide(run_dir)
    lines = []
    if promote:
        lines.append("Promoted to baseline (improved, nothing worse): "
                     + ", ".join(f"{c} {l}" for c, l in promote))
    if held:
        lines.append("Improved but NOT promoted: "
                     + "; ".join(f"{c} {l} ({why})" for c, l, why in held))
    if not lines:
        lines.append("No improvements to promote.")
    for line in lines:
        print(line)
    if dry_run or not promote:
        if not dry_run:
            (run_dir / "BaselineUpdate.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")
        return promote

    for comp, stem in COMPONENTS:
        locs = {l for c, l in promote if c == comp}
        if not locs:
            continue
        base_xlsx = baseline_dir / f"{stem}.xlsx"
        merge_workbook(base_xlsx, run_dir / f"{stem}.xlsx", locs)
        generate_html_report.main(input_xlsx=str(base_xlsx), output_html=str(baseline_dir / f"{stem}.html"),
                                  component=comp)

    (run_dir / "BaselineUpdate.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")
    with open(baseline_dir / "BASELINE_SOURCE.txt", "a", encoding="utf-8") as f:
        f.write(f"Improvements promoted from daily/{run_dir.name} "
                f"({datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC): "
                + ", ".join(f"{c} {l}" for c, l in promote) + "\n")
    return promote


def main():
    ap = argparse.ArgumentParser(description="Promote improved component/locales into the baseline")
    ap.add_argument("run_dir")
    ap.add_argument("--baseline", default=str(CLOUD_BASELINE_DIR))
    ap.add_argument("--dry-run", action="store_true", help="Only print what would be promoted")
    a = ap.parse_args()
    run(Path(a.run_dir), Path(a.baseline), a.dry_run)


if __name__ == "__main__":
    main()
