"""Build a short Slack summary of a daily run: what changed, where, and whether it matters.

Writes daily/<run>/slack_message.md (standard markdown, as the Slack connector expects) and prints it.
The cloud routine posts that file to #teamviewer-header-footer-compare with the Slack connector.

Three outcomes:
  - Run problem   : the check did not finish, a stage crashed, or a locale page errored
  - Changed       : DifferenceReport exists -> changes grouped across locales, worse first
  - No change     : one line

Usage: python3 scripts/slack_summary.py [--run-dir daily/<folder>] [--branch main]
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
REPO_URL = "https://github.com/rushi-patel-hztl/tv-header-footer-check"
BASELINE_SOURCE = ROOT / "baselines" / "cloud_baseline" / "BASELINE_SOURCE.txt"
REPORTS = [("Header", "HeaderCheck_Report.xlsx"), ("Footer", "FooterCheck_Report.xlsx")]
AEM, SITECORE = "Home page (AEM)", "Target page (Sitecore)"
MAX_CHARS = 3800  # Slack text element limit is 5000; leave room for the footer lines
LINK_CHANGES = {"Link added", "Link removed", "Link text changed", "Link URL changed"}


def sheet(path: Path, name: str) -> list[dict]:
    wb = load_workbook(path, read_only=True)
    if name not in wb.sheetnames:
        return []
    rows = list(wb[name].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:] if any(v is not None for v in r)]


def latest_run() -> Path:
    runs = sorted(p for p in (ROOT / "daily").iterdir() if p.is_dir())
    if not runs:
        sys.exit("No run folder found under daily/")
    return runs[-1]


def status_counts(run: Path) -> dict[str, dict[str, str]]:
    """{component: {locale: status}} from each report's Summary sheet."""
    out = {}
    for comp, name in REPORTS:
        if (run / name).exists():
            out[comp] = {r["Locale"]: r["Status"] for r in sheet(run / name, "Summary") if r.get("Locale")}
    return out


def fmt_locales(locs: list[str], all_locs: list[str]) -> str:
    locs = [l for l in all_locs if l in locs] + sorted(set(locs) - set(all_locs))
    if all_locs and len(locs) == len(all_locs):
        return f"all {len(locs)} locales"
    missing = [l for l in all_locs if l not in locs]
    if all_locs and len(locs) >= 12:
        return f"{len(locs)} locales (all except {', '.join(missing)})"
    return f"{len(locs)} locale{'s' if len(locs) > 1 else ''}: {', '.join(locs)}"


def norm_path(path: str | None) -> str:
    """Drop the locale prefix so the same link groups across locales: /fr/global/x -> /global/x."""
    if not path:
        return ""
    if path.startswith("tel:"):
        return "tel:"
    return re.sub(r"^/[a-z]{2}(-[a-z]+)?(?=/)", "", path.split("?")[0]).rstrip("/") or "/"


def group_changes(details: list[dict]) -> list[dict]:
    """Collapse per-locale Details rows into one line per identical change."""
    # Effect on the AEM vs Sitecore comparison, keyed by component/locale/label
    effect = {}
    for r in details:
        if r["Page"] == "Comparison":
            e = "worse" if "(worse)" in r["Change"] else "better" if "(better)" in r["Change"] else None
            effect[(r["Component"], r["Locale"], r["Link / Text"])] = e

    # Labels already covered by a link change: their Text added/removed rows are noise
    link_labels = set()
    for r in details:
        if r["Change"] in LINK_CHANGES:
            for lbl in (r["Link / Text"], r.get("Baseline"), r.get("Current")):
                link_labels.add((r["Component"], r["Locale"], r["Page"], lbl))

    groups: dict[tuple, dict] = {}
    for r in details:
        page, change = r["Page"], r["Change"]
        if page not in (AEM, SITECORE):
            continue
        comp, loc, label = r["Component"], r["Locale"], r["Link / Text"]
        if change in ("Text added", "Text removed") and (comp, loc, page, label) in link_labels:
            continue
        path = norm_path(r.get("Path"))
        if change in ("Text added", "Text removed"):
            key = (comp, page, change, label)
        elif change == "Link text changed":
            key = (comp, page, change, path, r.get("Current"))
        else:
            key = (comp, page, change, path)
        g = groups.setdefault(key, {"comp": comp, "page": page, "change": change, "path": path,
                                    "locales": [], "labels": {}, "old": {}, "effects": {}})
        g["locales"].append(loc)
        g["labels"][loc] = label
        g["old"][loc] = r.get("Baseline")
        g["effects"][loc] = effect.get((comp, loc, label)) or effect.get((comp, loc, r.get("Baseline")))
    return list(groups.values())


def describe(g: dict, all_locs: list[str]) -> str:
    where = "Sitecore page" if g["page"] == SITECORE else "AEM home page"
    labels = list(dict.fromkeys(g["labels"].values()))
    sample = labels[0] if labels else ""
    ch, path = g["change"], g["path"]
    if path == "tel:":
        nums = ", ".join(f"`{n}`" for n in labels[:3])
        what = f"Phone number link {ch.split()[1]} ({nums})"
    elif ch == "Link text changed":
        olds = list(dict.fromkeys(v for v in g["old"].values() if v))
        was = f"was localised, e.g. \"{olds[0]}\"" if len(olds) > 1 else f"was \"{olds[0]}\"" if olds else ""
        what = f"Link label changed to \"{sample}\" ({was}) on `{path}`"
    elif ch in ("Text added", "Text removed"):
        what = f"Text \"{sample}\" {ch.split()[1]}"
    else:
        eg = f"\"{sample}\"" + (" and translations" if len(labels) > 1 else "")
        what = f"{ch} {eg} → `{path}`"
    if g["page"] == AEM and effect_of(g) == "worse" and ch in ("Link added", "Link removed"):
        what += " (Sitecore page still has it)" if ch == "Link removed" else " (Sitecore page does not have it)"
    line = f"**{g['comp']}** · {where}: {what} — {fmt_locales(g['locales'], all_locs)}"
    if path == "tel:":
        line += "\n      _Phone numbers vary by visitor location; confirm before raising._"
    return line


def effect_of(g: dict) -> str:
    vals = {v for v in g["effects"].values() if v}
    if vals == {"worse"}:
        return "worse"
    if vals == {"better"}:
        return "better"
    return "mixed" if vals else "neutral"


def build(run: Path, branch: str) -> str:
    log = (run / "run.log").read_text(errors="ignore") if (run / "run.log").exists() else ""
    run_url = f"{REPO_URL}/tree/{branch}/daily/{run.name}"
    blob_url = f"{REPO_URL}/blob/{branch}/daily/{run.name}"
    when = re.sub(r"^(\d{4}-\d{2}-\d{2})_(\d{2})(\d{2})\d{2}$", r"\1 \2:\3 UTC", run.name)
    statuses = status_counts(run)
    all_locs = list(next(iter(statuses.values()), {}).keys())
    m = re.search(r"Copied from: daily/(\S+)", BASELINE_SOURCE.read_text()) if BASELINE_SOURCE.exists() else None
    baseline = m.group(1) if m else "cloud baseline"

    def status_line() -> str:
        parts = []
        for comp, st in statuses.items():
            c = {s: list(st.values()).count(s) for s in ("ERROR", "FAIL", "WARN", "PASS")}
            parts.append(f"{comp}: " + " · ".join(f"{n} {s}" for s, n in c.items() if n))
        return "  |  ".join(parts)

    links = f"[Run folder]({run_url}) · [run.log]({blob_url}/run.log)"

    # 1. Run problems: say so first, results may be incomplete
    problems = []
    if "=== Cloud run finished" not in log:
        problems.append("The run did not finish (no finish line in run.log).")
    for stage in ("HEADER CHECK", "FOOTER CHECK", "BASELINE COMPARE"):
        if f"{stage} FAILED" in log:
            err = re.findall(rf"{stage} FAILED:.*?\n(?:.*\n)*?([\w.]*(?:Error|Exception)[^\n]*)", log)
            problems.append(f"{stage.title()} crashed" + (f": `{err[-1][:200]}`" if err else "."))
    for comp, st in statuses.items():
        bad = [l for l, s in st.items() if s == "ERROR"]
        if bad:
            problems.append(f"{comp} page could not be checked for {', '.join(bad)} (status ERROR).")
    for what in ("Header element not found", "Footer element not found"):
        n = log.count(what)
        if n:
            problems.append(f"\"{what}\" × {n}. A page may be blocked or have changed structure.")
    if problems:
        body = "\n".join(f"• {p}" for p in problems)
        return (f":x: **Header/footer check had problems** — run {when}\n{body}\n"
                f"Results may be incomplete. **Do not read this as \"no change\".**\n"
                + (f"Status vs AEM: {status_line()}\n" if statuses else "") + links)

    diff = run / "DifferenceReport.xlsx"
    if not diff.exists():
        return (f":large_green_circle: **No change in header/footer** — run {when}\n"
                f"All {len(all_locs)} locales match the baseline ({baseline}) for header and footer.\n"
                f"Status vs AEM (unchanged): {status_line()}\n{links}")

    # 2. Changes vs baseline
    overview = sheet(diff, "Overview")
    groups = group_changes(sheet(diff, "Details"))
    order = {"worse": 0, "mixed": 1, "neutral": 2, "better": 3}
    groups.sort(key=lambda g: (order[effect_of(g)], -len(g["locales"]), g["comp"]))
    n_worse = sum(1 for g in groups if effect_of(g) in ("worse", "mixed"))
    dirs = [r.get("Direction") or "" for r in overview]
    pct = ", ".join(f"{dirs.count(d)} {d.split(' (')[0].lower()}" for d in dict.fromkeys(dirs) if d)
    locs_changed = sorted({r["Locale"] for r in overview})

    head = (f":red_circle: **Header/footer changed vs baseline** — run {when}\n"
            f"{len(groups)} distinct change{'s' if len(groups) != 1 else ''} across {len(locs_changed)} locales "
            f"({n_worse} move Sitecore further from AEM). Baseline: {baseline}.\n"
            f"Match % by locale/component: {pct}.\n")

    rank = {"ERROR": 0, "FAIL": 1, "WARN": 2, "PASS": 3}
    flips = [(rank.get(r["Current Status"], 0) - rank.get(r["Baseline Status"], 0),
              f"{r['Component']} {r['Locale']}: {r['Baseline Status']} → **{r['Current Status']}**")
             for r in overview if r.get("Baseline Status") != r.get("Current Status")]
    sections = []
    down = [f for d, f in flips if d < 0]
    up = [f for d, f in flips if d > 0]
    if down:
        sections.append(":rotating_light: **Status got worse**\n" + "\n".join(f"• {f}" for f in down))
    if up:
        sections.append(":arrow_up: **Status improved**\n" + "\n".join(f"• {f}" for f in up))
    titles = {"worse": ":warning: **Needs a look: Sitecore now further from AEM**",
              "mixed": ":grey_question: **Mixed effect**",
              "neutral": ":information_source: **Other changes (check they are intended)**",
              "better": ":white_check_mark: **Sitecore now closer to AEM (confirm intended)**"}
    for eff in order:
        items = [g for g in groups if effect_of(g) == eff]
        if items:
            sections.append(titles[eff] + "\n" + "\n".join(f"• {describe(g, all_locs)}" for g in items))

    text = head
    for i, s in enumerate(sections):
        if len(text) + len(s) > MAX_CHARS:
            text += f"\n_…{len(sections) - i} more section(s) in the report._\n"
            break
        text += "\n" + s + "\n"
    text += f"\nStatus vs AEM now: {status_line()}\n"
    text += f"Details: [DifferenceReport.html]({blob_url}/DifferenceReport.html) · {links}"
    return text


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--run-dir")
    ap.add_argument("--branch", default="main", help="Branch the run folder is pushed to (for links)")
    a = ap.parse_args()
    run = Path(a.run_dir) if a.run_dir else latest_run()
    text = build(run, a.branch)
    (run / "slack_message.md").write_text(text + "\n", encoding="utf-8")
    print(text)


if __name__ == "__main__":
    main()
