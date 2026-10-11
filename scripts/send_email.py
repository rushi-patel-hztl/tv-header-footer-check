"""Email the result of the latest daily run via Resend.

Env: RESEND_API_KEY (required), MAIL_TO (comma separated, required),
     MAIL_FROM (default "onboarding@resend.dev").
Usage: python3 scripts/send_email.py [--dry-run] [--run-dir daily/<folder>]
"""
import argparse
import base64
import html
import json
import os
import sys
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
GREEN, RED = "#1a7f37", "#cf222e"


def latest_run() -> Path:
    runs = sorted(p for p in (ROOT / "daily").iterdir() if p.is_dir())
    if not runs:
        sys.exit("No run folder found under daily/")
    return runs[-1]


def overview_rows(xlsx: Path) -> tuple[list, list]:
    ws = load_workbook(xlsx, read_only=True)["Overview"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return rows[0], rows[1:]


def overall_by_locale(xlsx: Path) -> tuple[dict, dict]:
    ws = load_workbook(xlsx, read_only=True)["Summary"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    col = {h: i for i, h in enumerate(rows[0])}
    pct, status = {}, {}
    for r in rows[1:]:
        if r[col["Locale"]]:
            pct[r[col["Locale"]]] = r[col["Overall Match %"]]
            status[r[col["Locale"]]] = r[col["Status"]]
    return pct, status


def changed_locales(run: Path, component: str) -> set:
    """Locales that differ from the baseline (any %, link or text change, error, added/removed)."""
    diff = run / "DifferenceReport.xlsx"
    if not diff.exists():
        return set()
    _, rows = overview_rows(diff)
    return {r[1] for r in rows if r[0] == component}


def summary_table(run: Path, name: str, title: str, component: str) -> str:
    cur, _ = overall_by_locale(run / name)
    changed = changed_locales(run, component)
    status = {loc: ("FAIL" if loc in changed else "PASS") for loc in cur}
    base_file = ROOT / "baselines" / "cloud_baseline" / name
    base = overall_by_locale(base_file)[0] if base_file.exists() else {}
    cell = "border:1px solid #ccc;padding:4px 8px"
    colors = {"PASS": GREEN, "FAIL": RED}
    tr = ""
    for loc, v in cur.items():
        b = base.get(loc)
        if b is None or v is None:
            delta, dc = "n/a", "#000"
        else:
            d = round(v - b, 2)
            delta, dc = (f"{d:+.2f}" if d else "0.00"), (GREEN if d > 0 else RED if d < 0 else "#000")
        tr += (
            f"<tr><td style='{cell}'>{html.escape(str(loc))}</td>"
            f"<td style='{cell};text-align:right'>{'' if b is None else f'{b}%'}</td>"
            f"<td style='{cell};text-align:right'>{v}%</td>"
            f"<td style='{cell};text-align:right;color:{dc}'>{delta}</td>"
            f"<td style='{cell};color:{colors.get(status[loc], '#000')}'>{html.escape(str(status[loc]))}</td></tr>"
        )
    th = "".join(f"<th style='{cell};background:#f0f0f0'>{h}</th>" for h in ("Locale", "Baseline %", "Last run %", "Delta", "vs baseline"))
    return f"<h3 style='margin:16px 0 6px'>{title}</h3><table style='border-collapse:collapse;font-size:13px'><tr>{th}</tr>{tr}</table>"


def tables(run: Path) -> str:
    return summary_table(run, "HeaderCheck_Report.xlsx", "Header - overall match % by locale", "Header") + summary_table(
        run, "FooterCheck_Report.xlsx", "Footer - overall match % by locale", "Footer"
    )


def build(run: Path) -> tuple[str, str, dict | None]:
    name = run.name
    finished = "=== Cloud run finished" in (run / "run.log").read_text(errors="ignore")
    diff = run / "DifferenceReport.xlsx"
    if not finished:
        line, color = "Check failed - run did not finish", RED
        body = f"<p>See <code>daily/{name}/run.log</code>.</p>"
        return f"[RED] TeamViewer header/footer check FAILED - {name}", wrap(line, color, name, body), None
    if not diff.exists():
        line, color = "No change from baseline", GREEN
        body = "<p>All 17 locales are identical to the cloud baseline for header and footer.</p>" + tables(run)
        return f"[GREEN] TeamViewer header/footer: no change from baseline - {name}", wrap(line, color, name, body), None
    headers, rows = overview_rows(diff)
    th = "".join(f"<th style='border:1px solid #ccc;padding:4px 8px;background:#f0f0f0'>{html.escape(str(h))}</th>" for h in headers)
    tr = "".join(
        "<tr>" + "".join(f"<td style='border:1px solid #ccc;padding:4px 8px'>{html.escape('' if c is None else str(c))}</td>" for c in r) + "</tr>"
        for r in rows
    )
    body = f"<p>{len(rows)} changed component/locale rows.</p><table style='border-collapse:collapse;font-size:13px'><tr>{th}</tr>{tr}</table><p>Full details: attached DifferenceReport.xlsx.</p>" + tables(run)
    att = {"filename": "DifferenceReport.xlsx", "content": base64.b64encode(diff.read_bytes()).decode()}
    return f"[RED] TeamViewer header/footer changed from baseline - {name}", wrap("Header/footer has changed from baseline", RED, name, body), att


def wrap(line: str, color: str, name: str, body: str) -> str:
    return (
        f"<div style='font-family:Arial,sans-serif'><h2 style='color:{color};margin:0 0 8px'>{html.escape(line)}</h2>"
        f"<p style='color:#555'>Run: {html.escape(name)} (UTC)</p>{body}</div>"
    )


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--run-dir")
    a = ap.parse_args()
    run = Path(a.run_dir) if a.run_dir else latest_run()
    subject, html_body, att = build(run)
    if a.dry_run:
        print(subject, "\n", html_body, "\nattachment:", bool(att))
        return
    key, to = os.environ.get("RESEND_API_KEY"), os.environ.get("MAIL_TO")
    if not key or not to:
        sys.exit("RESEND_API_KEY and MAIL_TO must be set")
    payload = {
        "from": os.environ.get("MAIL_FROM", "onboarding@resend.dev"),
        "to": [t.strip() for t in to.split(",")],
        "subject": subject,
        "html": html_body,
    }
    if att:
        payload["attachments"] = [att]
    req = urllib.request.Request(
        "https://api.resend.com/emails",
        data=json.dumps(payload).encode(),
        headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json", "User-Agent": "tv-check/1.0"},
    )
    with urllib.request.urlopen(req, timeout=30) as r:
        print("Email sent:", r.status, r.read().decode())


if __name__ == "__main__":
    main()
