"""
Compare header text and links between source and target TeamViewer pages.
Reads pairs from HeaderCheck.xlsx and writes a detailed comparison report.
"""

from __future__ import annotations

import re
import time
from collections import OrderedDict
from urllib.parse import urlparse, urlunparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright

INPUT_XLSX = r"c:\Users\rpatel\Desktop\Teamviewer\HeaderCheck.xlsx"
# New run uses separate files so prior reports are not overwritten
OUTPUT_XLSX = r"c:\Users\rpatel\Desktop\Teamviewer\HeaderCheck_Report_v8.xlsx"
OUTPUT_HTML = r"c:\Users\rpatel\Desktop\Teamviewer\HeaderCheck_Report_v8.html"

# Live site host used for both homepage (source) and KB page (target)
LIVE_HOST = "www.teamviewer.com"

# Domains to treat as equivalent when comparing link paths
HOST_ALIASES = {
    "www.teamviewer.com",
    "prod-sitecore.teamviewer.com",
    "teamviewer.com",
}


def locale_home_url(url: str) -> str:
    """Map any locale URL to that locale's homepage on the live site.

    e.g. .../de/global/support/knowledge-base -> https://www.teamviewer.com/de/
    """
    parsed = urlparse(str(url).strip())
    parts = [p for p in parsed.path.split("/") if p]
    locale = parts[0] if parts else ""
    path = f"/{locale}/" if locale else "/"
    return urlunparse(("https", LIVE_HOST, path, "", "", ""))


def live_kb_url(url: str) -> str:
    """Keep the KB path but force the live teamviewer.com host.

    e.g. prod-sitecore.../de/.../knowledge-base -> https://www.teamviewer.com/de/.../knowledge-base
    """
    parsed = urlparse(str(url).strip())
    path = parsed.path or "/"
    return urlunparse(("https", LIVE_HOST, path, "", "", ""))


def load_pairs(path: str) -> list[tuple[str, str]]:
    """Load Excel pairs, remapped to homepage (source) vs live KB (target)."""
    wb = load_workbook(path)
    ws = wb.active
    pairs: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        source, target = row[0], row[1]
        if not source or not target:
            continue
        # Source = locale homepage; target = same KB path on www.teamviewer.com
        key = (locale_home_url(source), live_kb_url(target))
        if key in seen:
            continue
        seen.add(key)
        pairs.append(key)
    return pairs


# Known locale aliases (source slug -> canonical)
LOCALE_ALIASES = {
    "ru-cis": "ru",
    "zh-cn": "cn",
    "pt-br": "pt",
}

# Known path slug renames between www (AEM) and Sitecore (applied after locale strip)
PATH_SLUG_ALIASES = [
    (r"/integrations\b", "/integration"),
    (r"/sport-partnership\b", "/sports-partnerships"),
    (r"/resources/reports\b", "/resources/report"),
    (r"/resources/trust-center\b", "/resources/trust-centers"),
    (r"/solutions/use-cases/nonprofit-education\b", "/solutions/use-cases/non-profit-education"),
    (r"/products/add-ons/remote-management\b", "/products/remote/add-ons/remote-management"),
    (r"/products/remote/pricing\b", "/pricing"),
]


def strip_locale(path: str) -> tuple[str, str]:
    """Return (locale_or_empty, path_without_locale)."""
    parts = [p for p in path.split("/") if p]
    if not parts:
        return "", "/"
    locale = parts[0].lower()
    locale = LOCALE_ALIASES.get(locale, locale)
    rest = "/" + "/".join(parts[1:]) if len(parts) > 1 else "/"
    return locale, rest


def apply_locale_aliases(path: str) -> str:
    """Rewrite locale slug aliases in-path (e.g. /ru-cis/... -> /ru/...).

    Keeps the rest of the path so ru-cis vs ru is treated as the same link,
    without ignoring the full locale segment.
    """
    parts = [p for p in path.split("/") if p]
    if not parts:
        return path or "/"
    locale = parts[0].lower()
    if locale in LOCALE_ALIASES:
        parts[0] = LOCALE_ALIASES[locale]
        return "/" + "/".join(parts)
    return path


def apply_slug_aliases(path: str) -> str:
    out = path
    for pattern, repl in PATH_SLUG_ALIASES:
        out = re.sub(pattern, repl, out, flags=re.I)
    return out


def normalize_href(href: str, *, structural: bool = False, keep_host: bool = False) -> str:
    """Normalize href for comparison.

    structural=False: path with locale aliases applied (ru-cis→ru); host stripped for TV domains
      unless keep_host=True (baseline regression — domain change is a real issue)
    structural=True: locale-stripped + known slug aliases + tel digits only
    keep_host=True: retain host (www. stripped only) so www vs prod-sitecore etc. do not match
    """
    if not href:
        return ""
    href = href.strip()

    # tel: compare digits only when structural
    if re.match(r"^tel:", href, re.I):
        if structural:
            digits = re.sub(r"\D", "", href)
            return f"tel:{digits}"
        return href.lower().rstrip("/")

    if re.match(r"^(mailto:|javascript:|#)", href, re.I):
        return href.lower().rstrip("/")

    parsed = urlparse(href)
    host = (parsed.netloc or "").lower().removeprefix("www.")
    path = parsed.path or "/"
    if path != "/" and path.endswith("/"):
        path = path[:-1]
    query = f"?{parsed.query}" if parsed.query else ""

    external_hosts = {
        "web.teamviewer.com",
        "login.teamviewer.com",
        "support.teamviewer.com",
        "service.teamviewer.com",
    }

    is_tv_site = "teamviewer.com" in host and host not in external_hosts
    if is_tv_site or not host:
        path_l = path.lower()
        if structural:
            _locale, rest = strip_locale(path_l)
            rest = apply_slug_aliases(rest)
            path_part = f"{rest}{query}"
        else:
            # Exact mode still aliases locale slugs (ru-cis == ru)
            path_part = f"{apply_locale_aliases(path_l)}{query}"
        # Baseline verify: keep host so domain drift fails the match
        if keep_host and host:
            return f"{host}{path_part}"
        return path_part

    if host in external_hosts:
        path_l = path.lower()
        if structural:
            return f"{host}{path_l}{query}".rstrip("/")
        return f"{host}{path_l}{query}"

    return f"{host}{path.lower()}{query}"

def clean_text(text: str) -> str:
    t = re.sub(r"\s+", " ", (text or "").strip())
    # Trailing punctuation often differs across CMS implementations
    t = re.sub(r"[.!?…]+$", "", t).strip()
    return t


def text_key(text: str) -> str:
    """Case-insensitive comparison key with light punctuation normalization."""
    t = clean_text(text).lower()
    # Drop surrounding quotes / odd separators
    t = t.strip("\"'`")
    return t


EXTRACT_HEADER_JS = """
() => {
  const dismissSelectors = [
    '#onetrust-accept-btn-handler',
    'button#onetrust-accept-btn-handler',
    '[aria-label="Accept All"]',
    'button[aria-label*="Accept"]',
    '.ot-sdk-container button[title*="Accept"]'
  ];
  for (const sel of dismissSelectors) {
    const btn = document.querySelector(sel);
    if (btn) { try { btn.click(); } catch (e) {} }
  }

  const header =
    document.querySelector('header') ||
    document.querySelector('[role="banner"]') ||
    document.querySelector('nav.navbar, nav[class*="header"], .experiencefragment header');

  if (!header) {
    return { error: 'Header element not found', links: [], texts: [], rawText: '' };
  }

  // Context nav (secondary nav bar, Sitecore only) is not part of the comparison:
  // the home page (AEM) has none. Remove it and all its children before extracting.
  header.querySelectorAll('[data-name="header-contextnav"], [data-track-nav="contextnav"]')
    .forEach((n) => n.remove());

  // Direct text nodes only — avoids parent aggregation differences:
  // Source: <span title> + <span desc>  (innerText has spaces)
  // Target: <p title><p desc>           (anchor innerText often concatenates with NO space)
  const directText = (el) => {
    if (!el) return '';
    return [...el.childNodes]
      .filter((n) => n.nodeType === 3)
      .map((n) => (n.textContent || '').trim())
      .filter(Boolean)
      .join(' ')
      .replace(/\\s+/g, ' ')
      .trim();
  };

  const normalizeLabel = (t) => (t || '').replace(/\\s+/g, ' ').trim();

  // Skip screen-reader-only nodes (source uses .utils-visually-hidden)
  // Do NOT use [aria-hidden=true] — closed mega-menus use that and would drop real labels.
  const isVisuallyHidden = (el) => {
    if (!el || !el.closest) return false;
    return !!el.closest('.utils-visually-hidden, .sr-only, .visually-hidden');
  };

  const pickLinkLabel = (a) => {
    // Visible text only — never use aria-label; skip visually-hidden nodes
    const titleSel = [
      '.cmp-menuitem__page-title',
      '.cmp-menuitem__title',
      '.cmp-menuitem__headline',
      '[data-name=\"title\"]',
      'h2', 'h3', 'h4',
      'p.font-weight-bold',
      'p.font-bold'
    ].join(',');
    const titleEl = a.querySelector(titleSel);
    if (titleEl && !isVisuallyHidden(titleEl)) {
      const t = normalizeLabel(directText(titleEl) || titleEl.innerText || '');
      if (t) return t.slice(0, 120);
    }
    // Leaf p/span/b/strong with direct visible text inside the anchor
    for (const child of a.querySelectorAll('p, span, strong, b')) {
      if (isVisuallyHidden(child)) continue;
      const t = normalizeLabel(directText(child));
      if (t && t.length >= 2 && t.length <= 80) return t;
    }
    // Build visible-only label: clone and strip visually-hidden nodes
    const clone = a.cloneNode(true);
    clone.querySelectorAll('.utils-visually-hidden, .sr-only, .visually-hidden').forEach((n) => n.remove());
    return normalizeLabel(clone.innerText || clone.textContent || '').slice(0, 120);
  };

  const links = [];
  const seen = new Set();
  for (const a of header.querySelectorAll('a[href]')) {
    const href = a.href || a.getAttribute('href') || '';
    if (!href || href.startsWith('javascript:')) continue;
    const text = pickLinkLabel(a);
    const targetAttr = (a.getAttribute('target') || '').trim();
    const key = href + '||' + text + '||' + targetAttr.toLowerCase();
    if (seen.has(key)) continue;
    seen.add(key);
    links.push({ text, href, target: targetAttr });
  }

  // Text labels: ONLY leaf/direct texts so DOM structure differences
  // (span+span vs p+p / h2+p) do not create false missing/additional labels.
  const texts = [];
  const textSeen = new Set();
  const addText = (raw) => {
    const t = normalizeLabel(raw);
    if (!t || t.length < 2 || t.length > 80) return;
    // Skip phone-only labels
    if (/^[+\\d\\s()-]{6,}$/.test(t)) return;
    const key = t.toLowerCase();
    if (textSeen.has(key)) return;
    textSeen.add(key);
    texts.push(t);
  };

  const collectLeaf = (el) => {
    if (!el || isVisuallyHidden(el)) return;
    const own = directText(el);
    if (own) addText(own);
  };

  // 1) Known AEM menuitem LEAF classes only (not headline wrappers that nest spans)
  header.querySelectorAll(
    '.cmp-menuitem__page-title, .cmp-menuitem__page-description, ' +
    '.cmp-menuitem__title, .cmp-menuitem__tagline, ' +
    '.cmp-navigation__item-title'
  ).forEach(collectLeaf);

  // 2) Leaf heading / paragraph / span / button / bold nodes with OWN text nodes
  //    Covers Sitecore <p>title</p><p>desc</p>, AEM leaf spans, and promo titles in <b>/<strong>
  //    Skips .utils-visually-hidden (and similar) screen-reader-only nodes.
  header.querySelectorAll(
    'h2, h3, h4, p, span, button, b, strong, [role=\"menuitem\"], .cmp-promobanner__tagline, .cmp-promobanner__title'
  ).forEach(collectLeaf);

  // 3) Top-nav anchors that are text-only (visible direct text only; ignore aria-label)
  header.querySelectorAll('a').forEach(collectLeaf);

  return {
    error: null,
    links,
    texts,
    rawText: (header.innerText || '').trim().replace(/\\s+/g, ' ').slice(0, 5000),
    linkCount: links.length,
    textCount: texts.length
  };
}
"""


def extract_header(page, url: str) -> dict:
    last_err = None
    for attempt in range(1, 4):
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=90000)
            last_err = None
            break
        except Exception as e:
            last_err = e
            page.wait_for_timeout(1500 * attempt)
    if last_err:
        raise last_err

    # Allow header XF / mega-menu to hydrate
    page.wait_for_timeout(3000)
    try:
        page.locator("header").first.wait_for(state="attached", timeout=15000)
    except Exception:
        pass
    # Cookie banner
    for sel in [
        "#onetrust-accept-btn-handler",
        'button:has-text("Accept All")',
        'button:has-text("Accept all")',
        'button:has-text("Alle akzeptieren")',
        'button:has-text("Tout accepter")',
    ]:
        try:
            loc = page.locator(sel).first
            if loc.is_visible(timeout=1500):
                loc.click(timeout=2000)
                page.wait_for_timeout(500)
                break
        except Exception:
            continue
    data = page.evaluate(EXTRACT_HEADER_JS)
    data["final_url"] = page.url
    data["requested_url"] = url
    return data

def is_target_blank(value) -> bool:
    return (value or "").strip().lower() == "_blank"


def _index_links(links: list[dict], structural: bool) -> OrderedDict[str, dict]:
    indexed: OrderedDict[str, dict] = OrderedDict()
    for link in links or []:
        exact = normalize_href(link["href"], structural=False)
        struct = normalize_href(link["href"], structural=True)
        key = struct if structural else exact
        if not key:
            continue
        blank = is_target_blank(link.get("target"))
        if key in indexed:
            # If any duplicate normalized link opens in a new tab, keep that
            if blank:
                indexed[key]["target_blank"] = True
                indexed[key]["target_attr"] = "_blank"
            continue
        indexed[key] = {
            "text": clean_text(link.get("text") or ""),
            "href": link["href"],
            "exact": exact,
            "struct": struct,
            "target_attr": (link.get("target") or "").strip(),
            "target_blank": blank,
        }
    return indexed


def compare_headers(source: dict, target: dict) -> dict:
    # Exact path comparison (locale aliases applied, e.g. ru-cis→ru)
    src_exact = _index_links(source.get("links") or [], structural=False)
    tgt_exact = _index_links(target.get("links") or [], structural=False)

    # Structural comparison (locale stripped + known slug aliases)
    src_struct = _index_links(source.get("links") or [], structural=True)
    tgt_struct = _index_links(target.get("links") or [], structural=True)

    exact_match_keys = set(src_exact) & set(tgt_exact)
    exact_missing = set(src_exact) - set(tgt_exact)
    exact_additional = set(tgt_exact) - set(src_exact)

    struct_match_keys = set(src_struct) & set(tgt_struct)
    struct_missing = set(src_struct) - set(tgt_struct)
    struct_additional = set(tgt_struct) - set(src_struct)

    # Renamed links = exact-missing but still present structurally
    renamed_links = []
    for k in sorted(exact_missing):
        struct_key = src_exact[k]["struct"]
        if struct_key in tgt_struct:
            renamed_links.append(
                {
                    "struct": struct_key,
                    "source_text": src_exact[k]["text"],
                    "target_text": tgt_struct[struct_key]["text"],
                    "source_href": src_exact[k]["href"],
                    "target_href": tgt_struct[struct_key]["href"],
                    "source_exact": src_exact[k]["exact"],
                    "target_exact": tgt_struct[struct_key]["exact"],
                }
            )

    matching_links = [
        {
            "norm": k,
            "source_text": src_exact[k]["text"],
            "target_text": tgt_exact[k]["text"],
            "source_href": src_exact[k]["href"],
            "target_href": tgt_exact[k]["href"],
            "text_match": src_exact[k]["text"].lower() == tgt_exact[k]["text"].lower(),
        }
        for k in sorted(exact_match_keys)
    ]

    # Truly missing = not found even after structural normalization
    missing_links = [
        {
            "norm": src_struct[k]["exact"],
            "struct": k,
            "text": src_struct[k]["text"],
            "href": src_struct[k]["href"],
        }
        for k in sorted(struct_missing)
    ]
    additional_links = [
        {
            "norm": tgt_struct[k]["exact"],
            "struct": k,
            "text": tgt_struct[k]["text"],
            "href": tgt_struct[k]["href"],
        }
        for k in sorted(struct_additional)
    ]

    # target="_blank" comparison on structurally paired links
    target_blank_mismatches = []
    target_blank_matches = 0
    for k in sorted(struct_match_keys):
        src_blank = bool(src_struct[k].get("target_blank"))
        tgt_blank = bool(tgt_struct[k].get("target_blank"))
        if src_blank == tgt_blank:
            target_blank_matches += 1
            continue
        if src_blank and not tgt_blank:
            issue = "Missing target=_blank on target"
        else:
            issue = "Extra target=_blank on target"
        target_blank_mismatches.append(
            {
                "struct": k,
                "source_exact": src_struct[k]["exact"],
                "target_exact": tgt_struct[k]["exact"],
                "text": src_struct[k]["text"] or tgt_struct[k]["text"],
                "source_href": src_struct[k]["href"],
                "target_href": tgt_struct[k]["href"],
                "source_target": src_struct[k].get("target_attr") or "(none)",
                "target_target": tgt_struct[k].get("target_attr") or "(none)",
                "source_blank": "Yes" if src_blank else "No",
                "target_blank": "Yes" if tgt_blank else "No",
                "issue": issue,
            }
        )

    # Text label comparison — leaf texts only, punctuation-normalized
    src_texts = {
        text_key(t): clean_text(t) for t in (source.get("texts") or []) if clean_text(t)
    }
    tgt_texts = {
        text_key(t): clean_text(t) for t in (target.get("texts") or []) if clean_text(t)
    }
    matching_texts = sorted(src_texts[k] for k in (set(src_texts) & set(tgt_texts)))
    missing_texts = sorted(src_texts[k] for k in (set(src_texts) - set(tgt_texts)))
    additional_texts = sorted(tgt_texts[k] for k in (set(tgt_texts) - set(src_texts)))

    union_exact = set(src_exact) | set(tgt_exact)
    exact_link_match_pct = (
        (len(exact_match_keys) / len(union_exact) * 100) if union_exact else 100.0
    )
    exact_coverage_pct = (
        (len(exact_match_keys) / len(src_exact) * 100) if src_exact else 100.0
    )

    union_struct = set(src_struct) | set(tgt_struct)
    struct_link_match_pct = (
        (len(struct_match_keys) / len(union_struct) * 100) if union_struct else 100.0
    )
    struct_coverage_pct = (
        (len(struct_match_keys) / len(src_struct) * 100) if src_struct else 100.0
    )

    union_texts = set(src_texts) | set(tgt_texts)
    text_match_pct = (
        (len(set(src_texts) & set(tgt_texts)) / len(union_texts) * 100) if union_texts else 100.0
    )
    text_coverage_pct = (
        (len(set(src_texts) & set(tgt_texts)) / len(src_texts) * 100) if src_texts else 100.0
    )

    # Overall prioritizes structural link match + text match
    overall_pct = round((struct_link_match_pct + text_match_pct) / 2, 2)

    paired_for_blank = len(struct_match_keys)
    target_blank_match_pct = (
        round(target_blank_matches / paired_for_blank * 100, 2) if paired_for_blank else 100.0
    )

    return {
        "source_link_count": len(src_exact),
        "target_link_count": len(tgt_exact),
        "matching_link_count": len(exact_match_keys),
        "renamed_link_count": len(renamed_links),
        "missing_link_count": len(missing_links),
        "additional_link_count": len(additional_links),
        "exact_link_match_pct": round(exact_link_match_pct, 2),
        "exact_coverage_pct": round(exact_coverage_pct, 2),
        "link_match_pct": round(struct_link_match_pct, 2),
        "link_coverage_pct": round(struct_coverage_pct, 2),
        "source_text_count": len(src_texts),
        "target_text_count": len(tgt_texts),
        "matching_text_count": len(matching_texts),
        "missing_text_count": len(missing_texts),
        "additional_text_count": len(additional_texts),
        "text_match_pct": round(text_match_pct, 2),
        "text_coverage_pct": round(text_coverage_pct, 2),
        "overall_match_pct": overall_pct,
        "matching_links": matching_links,
        "renamed_links": renamed_links,
        "missing_links": missing_links,
        "additional_links": additional_links,
        "matching_texts": matching_texts,
        "missing_texts": missing_texts,
        "additional_texts": additional_texts,
        "exact_missing_count": len(exact_missing),
        "exact_additional_count": len(exact_additional),
        "target_blank_paired_count": paired_for_blank,
        "target_blank_match_count": target_blank_matches,
        "target_blank_mismatch_count": len(target_blank_mismatches),
        "target_blank_match_pct": target_blank_match_pct,
        "target_blank_mismatches": target_blank_mismatches,
    }

def style_header(ws, row=1):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin


def autofit(ws, max_width=60):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col:
            val = str(cell.value) if cell.value is not None else ""
            length = max(length, min(len(val), max_width))
        ws.column_dimensions[letter].width = max(12, length + 2)


def write_report(results: list[dict], output_path: str):
    wb = Workbook()

    # --- Summary ---
    ws = wb.active
    ws.title = "Summary"
    headers = [
        "#",
        "Locale",
        "Source URL",
        "Target URL",
        "Overall Match %",
        "Structural Link Match %",
        "Link Coverage % (src→tgt)",
        "Exact Link Match %",
        "Text Match %",
        "Text Coverage % (src→tgt)",
        "Src Links",
        "Tgt Links",
        "Exact Matching Links",
        "Renamed Slug Links",
        "Truly Missing in Target",
        "Truly Additional in Target",
        "Src Texts",
        "Tgt Texts",
        "Matching Texts",
        "Missing Texts",
        "Additional Texts",
        "target=_blank Paired",
        "target=_blank Matches",
        "target=_blank Mismatches",
        "target=_blank Match %",
        "Source Error",
        "Target Error",
        "Status",
    ]
    ws.append(headers)
    style_header(ws)

    green = PatternFill("solid", fgColor="C6EFCE")
    yellow = PatternFill("solid", fgColor="FFEB9C")
    red = PatternFill("solid", fgColor="FFC7CE")

    for i, r in enumerate(results, 1):
        cmp = r["comparison"]
        overall = cmp["overall_match_pct"]
        if r["source"].get("error") or r["target"].get("error"):
            status = "ERROR"
        elif overall >= 95 and cmp["missing_link_count"] == 0:
            status = "PASS"
        elif overall >= 80:
            status = "WARN"
        else:
            status = "FAIL"
        row = [
            i,
            r["locale"],
            r["source_url"],
            r["target_url"],
            overall,
            cmp["link_match_pct"],
            cmp["link_coverage_pct"],
            cmp["exact_link_match_pct"],
            cmp["text_match_pct"],
            cmp["text_coverage_pct"],
            cmp["source_link_count"],
            cmp["target_link_count"],
            cmp["matching_link_count"],
            cmp["renamed_link_count"],
            cmp["missing_link_count"],
            cmp["additional_link_count"],
            cmp["source_text_count"],
            cmp["target_text_count"],
            cmp["matching_text_count"],
            cmp["missing_text_count"],
            cmp["additional_text_count"],
            cmp["target_blank_paired_count"],
            cmp["target_blank_match_count"],
            cmp["target_blank_mismatch_count"],
            cmp["target_blank_match_pct"],
            r["source"].get("error") or "",
            r["target"].get("error") or "",
            status,
        ]
        ws.append(row)
        fill = green if status == "PASS" else yellow if status == "WARN" else red
        ws.cell(row=i + 1, column=5).fill = fill
        ws.cell(row=i + 1, column=28).fill = fill

    ws.freeze_panes = "A2"
    autofit(ws)

    # --- Matching Links (exact path) ---
    ws_m = wb.create_sheet("Matching Links")
    ws_m.append(
        [
            "Locale",
            "Exact Path",
            "Source Text",
            "Target Text",
            "Text Match",
            "Source Href",
            "Target Href",
        ]
    )
    style_header(ws_m)
    for r in results:
        for link in r["comparison"]["matching_links"]:
            ws_m.append(
                [
                    r["locale"],
                    link["norm"],
                    link["source_text"],
                    link["target_text"],
                    "Yes" if link["text_match"] else "No",
                    link["source_href"],
                    link["target_href"],
                ]
            )
    ws_m.freeze_panes = "A2"
    autofit(ws_m)

    # --- Renamed slug links ---
    ws_ren = wb.create_sheet("Renamed Slug Links")
    ws_ren.append(
        [
            "Locale",
            "Structural Key",
            "Source Exact Path",
            "Target Exact Path",
            "Source Text",
            "Target Text",
            "Source Href",
            "Target Href",
        ]
    )
    style_header(ws_ren)
    for r in results:
        for link in r["comparison"]["renamed_links"]:
            ws_ren.append(
                [
                    r["locale"],
                    link["struct"],
                    link["source_exact"],
                    link["target_exact"],
                    link["source_text"],
                    link["target_text"],
                    link["source_href"],
                    link["target_href"],
                ]
            )
    ws_ren.freeze_panes = "A2"
    autofit(ws_ren)

    # --- Truly Missing Links (in source, not in target even structurally) ---
    ws_miss = wb.create_sheet("Missing in Target")
    ws_miss.append(
        ["Locale", "Exact Path", "Structural Key", "Link Text", "Source Href"]
    )
    style_header(ws_miss)
    for r in results:
        for link in r["comparison"]["missing_links"]:
            ws_miss.append(
                [r["locale"], link["norm"], link["struct"], link["text"], link["href"]]
            )
    ws_miss.freeze_panes = "A2"
    autofit(ws_miss)

    # --- Truly Additional Links ---
    ws_add = wb.create_sheet("Additional in Target")
    ws_add.append(
        ["Locale", "Exact Path", "Structural Key", "Link Text", "Target Href"]
    )
    style_header(ws_add)
    for r in results:
        for link in r["comparison"]["additional_links"]:
            ws_add.append(
                [r["locale"], link["norm"], link["struct"], link["text"], link["href"]]
            )
    ws_add.freeze_panes = "A2"
    autofit(ws_add)

    # --- target="_blank" mismatches (structurally paired links) ---
    ws_blank = wb.create_sheet("target=_blank Mismatches")
    ws_blank.append(
        [
            "Locale",
            "Issue",
            "Structural Key",
            "Link Text",
            "Source target attr",
            "Target target attr",
            "Source _blank",
            "Target _blank",
            "Source Exact Path",
            "Target Exact Path",
            "Source Href",
            "Target Href",
        ]
    )
    style_header(ws_blank)
    for r in results:
        for link in r["comparison"]["target_blank_mismatches"]:
            ws_blank.append(
                [
                    r["locale"],
                    link["issue"],
                    link["struct"],
                    link["text"],
                    link["source_target"],
                    link["target_target"],
                    link["source_blank"],
                    link["target_blank"],
                    link["source_exact"],
                    link["target_exact"],
                    link["source_href"],
                    link["target_href"],
                ]
            )
    ws_blank.freeze_panes = "A2"
    autofit(ws_blank)

    # --- Matching Texts ---
    ws_mt = wb.create_sheet("Matching Texts")
    ws_mt.append(["Locale", "Matching Text Label"])
    style_header(ws_mt)
    for r in results:
        for t in r["comparison"]["matching_texts"]:
            ws_mt.append([r["locale"], t])
    ws_mt.freeze_panes = "A2"
    autofit(ws_mt)

    # --- Missing Texts ---
    ws_mxt = wb.create_sheet("Missing Texts")
    ws_mxt.append(["Locale", "Missing Text (in Source only)"])
    style_header(ws_mxt)
    for r in results:
        for t in r["comparison"]["missing_texts"]:
            ws_mxt.append([r["locale"], t])
    ws_mxt.freeze_panes = "A2"
    autofit(ws_mxt)

    # --- Additional Texts ---
    ws_at = wb.create_sheet("Additional Texts")
    ws_at.append(["Locale", "Additional Text (in Target only)"])
    style_header(ws_at)
    for r in results:
        for t in r["comparison"]["additional_texts"]:
            ws_at.append([r["locale"], t])
    ws_at.freeze_panes = "A2"
    autofit(ws_at)

    # --- Per-pair detail ---
    ws_d = wb.create_sheet("Pair Detail")
    ws_d.append(
        [
            "Locale",
            "Source Final URL",
            "Target Final URL",
            "Source Link Count (raw)",
            "Target Link Count (raw)",
            "Overall Match %",
            "Structural Link Match %",
            "Exact Link Match %",
            "Renamed Count",
            "Truly Missing Sample",
            "Truly Additional Sample",
            "Renamed Sample",
            "Missing Texts Sample",
            "Additional Texts Sample",
            "target=_blank Mismatch Count",
            "target=_blank Mismatch Sample",
        ]
    )
    style_header(ws_d)
    for r in results:
        cmp = r["comparison"]
        ws_d.append(
            [
                r["locale"],
                r["source"].get("final_url", ""),
                r["target"].get("final_url", ""),
                r["source"].get("linkCount", len(r["source"].get("links") or [])),
                r["target"].get("linkCount", len(r["target"].get("links") or [])),
                cmp["overall_match_pct"],
                cmp["link_match_pct"],
                cmp["exact_link_match_pct"],
                cmp["renamed_link_count"],
                " | ".join(f'{x["text"]} ({x["norm"]})' for x in cmp["missing_links"][:15]),
                " | ".join(f'{x["text"]} ({x["norm"]})' for x in cmp["additional_links"][:15]),
                " | ".join(
                    f'{x["source_exact"]} -> {x["target_exact"]}'
                    for x in cmp["renamed_links"][:15]
                ),
                " | ".join(cmp["missing_texts"][:20]),
                " | ".join(cmp["additional_texts"][:20]),
                cmp["target_blank_mismatch_count"],
                " | ".join(
                    f'{x["issue"]}: {x["struct"]}'
                    for x in cmp["target_blank_mismatches"][:15]
                ),
            ]
        )
    ws_d.freeze_panes = "A2"
    autofit(ws_d, max_width=80)

    # Legend / notes
    ws_n = wb.create_sheet("Notes")
    notes = [
        ["Field", "Meaning"],
        [
            "Exact Link Match %",
            "Paths match after locale aliases (ru-cis→ru). Other slug renames still count as Renamed.",
        ],
        [
            "Structural Link Match %",
            "Locale stripped + known slug aliases applied (integrations→integration, etc).",
        ],
        [
            "Renamed Slug Links",
            "Present on both sides after structural normalization but exact paths differ.",
        ],
        [
            "Missing in Target",
            "In source header but not found on target even after structural normalization.",
        ],
        [
            "Additional in Target",
            "In target header but not found on source even after structural normalization.",
        ],
        [
            "Overall Match %",
            "Average of Structural Link Match % and Text Match %.",
        ],
        [
            "Text Match %",
            "Jaccard similarity of unique short header text labels (case-insensitive).",
        ],
        [
            "Known aliases",
            "ru-cis→ru; /integrations→/integration; sport-partnership→sports-partnerships; reports→report; trust-center→trust-centers; nonprofit-education→non-profit-education; add-ons/remote-management path; remote/pricing→pricing; tel digits normalized.",
        ],
        [
            "target=_blank Mismatches",
            "Compared only on structurally paired links (normalized path). Reports when source has target=_blank and target does not, or vice versa.",
        ],
    ]
    for row in notes:
        ws_n.append(row)
    style_header(ws_n)
    autofit(ws_n, max_width=100)

    wb.save(output_path)

def locale_from_url(url: str) -> str:
    parts = urlparse(url).path.strip("/").split("/")
    return parts[0] if parts else "?"


def main():
    pairs = load_pairs(INPUT_XLSX)
    print(f"Loaded {len(pairs)} unique pairs")
    results = []

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        context = browser.new_context(
            viewport={"width": 1440, "height": 900},
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/122.0.0.0 Safari/537.36"
            ),
            ignore_https_errors=True,
        )
        page = context.new_page()
        page.set_default_timeout(60000)

        for idx, (source_url, target_url) in enumerate(pairs, 1):
            locale = locale_from_url(source_url)
            print(f"\n[{idx}/{len(pairs)}] Locale={locale}")
            print(f"  Source: {source_url}")
            src_data = {"error": None, "links": [], "texts": []}
            tgt_data = {"error": None, "links": [], "texts": []}
            try:
                src_data = extract_header(page, source_url)
                print(f"  Source links={src_data.get('linkCount')} texts={src_data.get('textCount')}")
            except Exception as e:
                src_data["error"] = str(e)
                print(f"  Source ERROR: {e}")

            print(f"  Target: {target_url}")
            try:
                tgt_data = extract_header(page, target_url)
                print(f"  Target links={tgt_data.get('linkCount')} texts={tgt_data.get('textCount')}")
            except Exception as e:
                tgt_data["error"] = str(e)
                print(f"  Target ERROR: {e}")

            comparison = compare_headers(src_data, tgt_data)
            print(
                f"  Overall={comparison['overall_match_pct']}% "
                f"StructLinks={comparison['link_match_pct']}% "
                f"ExactLinks={comparison['exact_link_match_pct']}% "
                f"Texts={comparison['text_match_pct']}% "
                f"Renamed={comparison['renamed_link_count']} "
                f"Missing={comparison['missing_link_count']} "
                f"Additional={comparison['additional_link_count']} "
                f"BlankMismatch={comparison['target_blank_mismatch_count']}"
            )
            results.append(
                {
                    "locale": locale,
                    "source_url": source_url,
                    "target_url": target_url,
                    "source": src_data,
                    "target": tgt_data,
                    "comparison": comparison,
                }
            )

        browser.close()

    write_report(results, OUTPUT_XLSX)
    print(f"\nReport saved to: {OUTPUT_XLSX}")
    try:
        from generate_html_report import main as write_html_report

        write_html_report(
            input_xlsx=OUTPUT_XLSX,
            output_html=OUTPUT_HTML,
            component="Header",
        )
        print(f"HTML report saved to: {OUTPUT_HTML}")
    except Exception as html_err:
        print(f"HTML report generation skipped: {html_err}")

    print("\n" + "=" * 120)
    print(
        f"{'Locale':<8} {'Overall%':>9} {'Struct%':>9} {'Exact%':>9} {'Text%':>8} "
        f"{'Rename':>7} {'Miss':>5} {'Add':>5} {'BlankMM':>7} {'Status':>8}"
    )
    print("=" * 120)
    for r in results:
        c = r["comparison"]
        overall = c["overall_match_pct"]
        status = (
            "ERROR"
            if r["source"].get("error") or r["target"].get("error")
            else "PASS"
            if overall >= 95 and c["missing_link_count"] == 0
            else "WARN"
            if overall >= 80
            else "FAIL"
        )
        print(
            f"{r['locale']:<8} {overall:>9.2f} {c['link_match_pct']:>9.2f} "
            f"{c['exact_link_match_pct']:>9.2f} {c['text_match_pct']:>8.2f} "
            f"{c['renamed_link_count']:>7} {c['missing_link_count']:>5} "
            f"{c['additional_link_count']:>5} {c['target_blank_mismatch_count']:>7} "
            f"{status:>8}"
        )
    print("=" * 120)

if __name__ == "__main__":
    main()
