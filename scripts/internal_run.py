"""
One-off runner: header + footer of www.teamviewer.com vs internal.teamviewer.com, same locale home page.

internal.teamviewer.com is behind Cloudflare Access, so this runs on a machine whose network Access
already lets through (not the cloud routine); no login step. Optionally set CF_ACCESS_CLIENT_ID and
CF_ACCESS_CLIENT_SECRET to send an Access service token from anywhere else.

Every page that still lands on the Cloudflare Access screen is reported as a page error, never as
a 0% match.

Writes internal_compare/<YYYY-MM-DD_HHMMSS>/ with:
  HeaderCheck_Report.xlsx / .html
  FooterCheck_Report.xlsx / .html
  run.log

Usage (from the repo root):
  python -u scripts/internal_run.py
  python -u scripts/internal_run.py --input other.xlsx --out-root some/folder
"""

from __future__ import annotations

import argparse
import os
import sys
import traceback
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright as _real_sync_playwright

SCRIPTS_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPTS_DIR.parent
sys.path.insert(0, str(SCRIPTS_DIR))

import footer_compare as fc  # noqa: E402
import header_compare as hc  # noqa: E402
import pacing  # noqa: E402

DEFAULT_INPUT = REPO_ROOT / "input" / "InternalCheck.xlsx"
DEFAULT_OUT_ROOT = REPO_ROOT / "internal_compare"
INTERNAL_HOST = "internal.teamviewer.com"


class _Tee:
    """Write to both the console and run.log."""

    def __init__(self, *streams):
        self.streams = streams

    def write(self, data):
        for s in self.streams:
            s.write(data)
        return len(data)

    def flush(self):
        for s in self.streams:
            s.flush()


def access_headers() -> dict:
    cid, secret = os.environ.get("CF_ACCESS_CLIENT_ID"), os.environ.get("CF_ACCESS_CLIENT_SECRET")
    if cid and secret:
        return {"CF-Access-Client-Id": cid, "CF-Access-Client-Secret": secret}
    return {}


def on_access_screen(page) -> bool:
    host = urlparse(page.url).netloc.lower()
    try:
        title = page.title()
    except Exception:
        title = ""
    return host.endswith("cloudflareaccess.com") or "Cloudflare Access" in title


class _AuthedPlaywright:
    """sync_playwright() stand-in whose browser contexts carry the Access service token, if set."""

    def __init__(self):
        self._cm = _real_sync_playwright()

    def __enter__(self):
        p = self._cm.__enter__()
        launch = p.chromium.launch

        def authed_launch(*args, **kwargs):
            browser = launch(*args, **kwargs)
            new_context = browser.new_context

            def authed_context(*a, **kw):
                headers = access_headers()
                if headers:
                    kw["extra_http_headers"] = {**kw.get("extra_http_headers", {}), **headers}
                return new_context(*a, **kw)

            browser.new_context = authed_context
            return browser

        p.chromium.launch = authed_launch
        return p

    def __exit__(self, *exc):
        return self._cm.__exit__(*exc)


def load_pairs(path: str) -> list[tuple[str, str]]:
    """Source/Target pairs as written (header_compare.load_pairs would force the www host)."""
    ws = load_workbook(path).active
    pairs: list[tuple[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] and row[1]:
            pairs.append((str(row[0]).strip(), str(row[1]).strip()))
    return pairs


def _guard_access(extract):
    """Report a page left on the Cloudflare Access screen as an error, not as an empty header/footer."""

    def wrapper(page, url):
        data = extract(page, url)
        if on_access_screen(page):
            data["error"] = "Blocked by Cloudflare Access - page not checked"
            data["links"], data["texts"] = [], []
        return data

    return wrapper


def main() -> int:
    ap = argparse.ArgumentParser(description="www vs internal.teamviewer.com header/footer check")
    ap.add_argument("--input", default=str(DEFAULT_INPUT), help="Source/Target URL pairs (xlsx)")
    ap.add_argument("--out-root", default=str(DEFAULT_OUT_ROOT), help="Folder for run folders")
    args = ap.parse_args()

    run_dir = Path(args.out_root) / datetime.now().strftime("%Y-%m-%d_%H%M%S")
    run_dir.mkdir(parents=True, exist_ok=False)
    log = open(run_dir / "run.log", "a", encoding="utf-8", buffering=1)
    sys.stdout = sys.stderr = _Tee(sys.__stdout__, log)
    print(f"=== www vs internal run started {datetime.now():%Y-%m-%d %H:%M:%S} ===")
    print(f"Run folder: {run_dir}")

    hc.sync_playwright = fc.sync_playwright = _AuthedPlaywright
    hc.load_pairs = fc.load_pairs = load_pairs
    hc.INPUT_XLSX = fc.INPUT_XLSX = args.input
    hc.OUTPUT_XLSX = str(run_dir / "HeaderCheck_Report.xlsx")
    hc.OUTPUT_HTML = str(run_dir / "HeaderCheck_Report.html")
    fc.OUTPUT_XLSX = str(run_dir / "FooterCheck_Report.xlsx")
    fc.OUTPUT_HTML = str(run_dir / "FooterCheck_Report.html")
    hc.extract_header = _guard_access(hc.extract_header)
    fc.extract_footer = _guard_access(fc.extract_footer)
    pacing.install(hc, fc)

    failures = 0
    for name, run in (("HEADER", hc.main), ("FOOTER", fc.main)):
        print(f"\n##### {name} CHECK #####")
        try:
            run()
        except Exception:
            failures += 1
            print(f"{name} CHECK FAILED:")
            traceback.print_exc()

    # The HTML subtitle names the daily check's hosts; name this run's
    for html in run_dir.glob("*.html"):
        text = html.read_text(encoding="utf-8")
        html.write_text(text.replace("Target (prod-sitecore.teamviewer.com)", f"Target ({INTERNAL_HOST})"),
                        encoding="utf-8")

    print(f"\n=== Run finished {datetime.now():%Y-%m-%d %H:%M:%S} (failures={failures}) ===")
    sys.stdout = sys.__stdout__
    sys.stderr = sys.__stderr__
    log.close()
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
