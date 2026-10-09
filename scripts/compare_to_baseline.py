"""
Compare a nightly run's Header/Footer reports against the fixed baseline.

If ANY match percentage changed for any locale (even 0.01%), or either page's
links / texts changed, writes a difference report into the run folder:
  DifferenceReport.xlsx / DifferenceReport.html

Changes are attributed to the page they happened on:
  Source = Home page   (AEM)
  Target = Target page (Sitecore)

Page link/text lists are rebuilt from each report's sheets:
  Home page   = Matching Links + Missing in Target + Renamed (source side)
  Target page = Matching Links + Additional in Target + Renamed (target side)

Usage:
  python compare_to_baseline.py <run_dir> [--baseline <baseline_dir>]
"""

from __future__ import annotations

import argparse
import html
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from header_compare import autofit, style_header

BASELINE_DIR = Path(r"c:\Users\rpatel\Desktop\Teamviewer\baselines\daily_baseline")

COMPONENTS = [("Header", "HeaderCheck_Report.xlsx"), ("Footer", "FooterCheck_Report.xlsx")]
SOURCE_PAGE = "Home page (AEM)"
TARGET_PAGE = "Target page (Sitecore)"

# Percentages that trigger a difference report when they change
PCT_FIELDS = [
    "Overall Match %",
    "Structural Link Match %",
    "Exact Link Match %",
    "Text Match %",
    "target=_blank Match %",
]
COUNT_FIELDS = [
    "Src Links",
    "Tgt Links",
    "Truly Missing in Target",
    "Truly Additional in Target",
    "Renamed Slug Links",
    "target=_blank Mismatches",
]
STATUS_RANK = {"ERROR": 0, "FAIL": 1, "WARN": 2, "PASS": 3}


# ---------------------------------------------------------------------------
# Load one report into per-locale page snapshots
# ---------------------------------------------------------------------------


def _rows(wb, sheet: str) -> list[dict]:
    if sheet not in wb.sheetnames:
        return []
    ws = wb[sheet]
    headers = [c.value for c in ws[1]]
    return [
        dict(zip(headers, r))
        for r in ws.iter_rows(min_row=2, values_only=True)
        if any(v is not None for v in r)
    ]


def load_report(path: Path) -> dict[str, dict]:
    """Return {locale: {summary, source_links, target_links, source_texts, target_texts, ...}}."""
    wb = load_workbook(path, read_only=True)
    data: dict[str, dict] = {}

    def loc(locale) -> dict:
        return data.setdefault(
            str(locale),
            {
                "summary": {},
                "source_links": {},  # exact path -> (text, href)
                "target_links": {},
                "source_texts": set(),
                "target_texts": set(),
                "missing": {},  # structural key -> text
                "additional": {},
                "blank": {},  # (issue, structural key) -> text
            },
        )

    for r in _rows(wb, "Summary"):
        loc(r["Locale"])["summary"] = r

    for r in _rows(wb, "Matching Links"):
        d = loc(r["Locale"])
        d["source_links"][r["Exact Path"]] = (r["Source Text"] or "", r["Source Href"] or "")
        d["target_links"][r["Exact Path"]] = (r["Target Text"] or "", r["Target Href"] or "")

    for r in _rows(wb, "Renamed Slug Links"):
        d = loc(r["Locale"])
        # setdefault: a path can also be a Matching Link on that page; keep that entry
        d["source_links"].setdefault(r["Source Exact Path"], (r["Source Text"] or "", r["Source Href"] or ""))
        d["target_links"].setdefault(r["Target Exact Path"], (r["Target Text"] or "", r["Target Href"] or ""))

    for r in _rows(wb, "Missing in Target"):
        d = loc(r["Locale"])
        d["source_links"].setdefault(r["Exact Path"], (r["Link Text"] or "", r["Source Href"] or ""))
        d["missing"][r["Structural Key"]] = r["Link Text"] or ""

    for r in _rows(wb, "Additional in Target"):
        d = loc(r["Locale"])
        d["target_links"].setdefault(r["Exact Path"], (r["Link Text"] or "", r["Target Href"] or ""))
        d["additional"][r["Structural Key"]] = r["Link Text"] or ""

    for r in _rows(wb, "target=_blank Mismatches"):
        loc(r["Locale"])["blank"][(r["Issue"], r["Structural Key"])] = r["Link Text"] or ""

    for r in _rows(wb, "Matching Texts"):
        d = loc(r["Locale"])
        d["source_texts"].add(r["Matching Text Label"])
        d["target_texts"].add(r["Matching Text Label"])
    for r in _rows(wb, "Missing Texts"):
        loc(r["Locale"])["source_texts"].add(r["Missing Text (in Source only)"])
    for r in _rows(wb, "Additional Texts"):
        loc(r["Locale"])["target_texts"].add(r["Additional Text (in Target only)"])

    wb.close()
    return data


# ---------------------------------------------------------------------------
# Diff
# ---------------------------------------------------------------------------


def _num(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def _page_changes(component, locale, page, base_links, cur_links, base_texts, cur_texts):
    out = []
    for path in sorted(set(cur_links) - set(base_links)):
        text, href = cur_links[path]
        out.append((component, locale, page, "Link added", text, path, "", href))
    for path in sorted(set(base_links) - set(cur_links)):
        text, href = base_links[path]
        out.append((component, locale, page, "Link removed", text, path, href, ""))
    for path in sorted(set(base_links) & set(cur_links)):
        old_t, cur_t = base_links[path][0], cur_links[path][0]
        if (old_t or "").strip() != (cur_t or "").strip():
            out.append((component, locale, page, "Link text changed", cur_t, path, old_t, cur_t))
        old_h, cur_h = base_links[path][1], cur_links[path][1]
        if old_h and cur_h and old_h != cur_h:
            out.append((component, locale, page, "Link URL changed", cur_t, path, old_h, cur_h))
    for t in sorted(cur_texts - base_texts):
        out.append((component, locale, page, "Text added", t, "", "", t))
    for t in sorted(base_texts - cur_texts):
        out.append((component, locale, page, "Text removed", t, "", t, ""))
    return out


def _comparison_changes(component, locale, b, c):
    """Effect on the page1 vs page2 comparison (what the score reacts to)."""
    out = []
    for k in sorted(set(c["missing"]) - set(b["missing"])):
        out.append((component, locale, "Comparison", "Newly missing in target (worse)", c["missing"][k], k, "", ""))
    for k in sorted(set(b["missing"]) - set(c["missing"])):
        out.append((component, locale, "Comparison", "No longer missing in target (better)", b["missing"][k], k, "", ""))
    for k in sorted(set(c["additional"]) - set(b["additional"])):
        out.append((component, locale, "Comparison", "Newly additional in target (worse)", c["additional"][k], k, "", ""))
    for k in sorted(set(b["additional"]) - set(c["additional"])):
        out.append((component, locale, "Comparison", "No longer additional in target (better)", b["additional"][k], k, "", ""))
    for k in sorted(set(c["blank"]) - set(b["blank"]), key=str):
        out.append((component, locale, "Comparison", f"New {k[0]} (worse)", c["blank"][k], k[1], "", ""))
    for k in sorted(set(b["blank"]) - set(c["blank"]), key=str):
        out.append((component, locale, "Comparison", f"Resolved {k[0]} (better)", b["blank"][k], k[1], "", ""))
    return out


def diff_component(component: str, base: dict, cur: dict):
    overview, details = [], []
    for locale in sorted(set(base) | set(cur), key=lambda x: list(base).index(x) if x in base else 999):
        b, c = base.get(locale), cur.get(locale)
        if b is None or c is None:
            overview.append(
                {
                    "Component": component,
                    "Locale": locale,
                    "Direction": "Locale added" if b is None else "Locale removed",
                    "Baseline Status": (b or {}).get("summary", {}).get("Status", "—") if b else "—",
                    "Current Status": c["summary"].get("Status", "—") if c else "—",
                }
            )
            continue

        bs, cs = b["summary"], c["summary"]
        pct_changes = {
            f: (_num(bs.get(f)), _num(cs.get(f)))
            for f in PCT_FIELDS
            if _num(bs.get(f)) != _num(cs.get(f))
        }
        status_changed = bs.get("Status") != cs.get("Status")
        errors = [e for e in (cs.get("Source Error"), cs.get("Target Error")) if e]

        rows = []
        rows += _page_changes(component, locale, SOURCE_PAGE, b["source_links"], c["source_links"],
                              b["source_texts"], c["source_texts"])
        rows += _page_changes(component, locale, TARGET_PAGE, b["target_links"], c["target_links"],
                              b["target_texts"], c["target_texts"])
        rows += _comparison_changes(component, locale, b, c)
        for e in errors:
            rows.append((component, locale, "Run", "Page error", e, "", "", ""))

        if not (pct_changes or status_changed or rows):
            continue

        b_over, c_over = _num(bs.get("Overall Match %")), _num(cs.get("Overall Match %"))
        if errors:
            direction = "Error"
        elif b_over is not None and c_over is not None and c_over > b_over:
            direction = "Improved"
        elif b_over is not None and c_over is not None and c_over < b_over:
            direction = "Regressed"
        elif STATUS_RANK.get(cs.get("Status"), 0) != STATUS_RANK.get(bs.get("Status"), 0):
            direction = "Improved" if STATUS_RANK[cs["Status"]] > STATUS_RANK[bs["Status"]] else "Regressed"
        else:
            direction = "Changed (overall % same)"

        row = {
            "Component": component,
            "Locale": locale,
            "Direction": direction,
            "Baseline Status": bs.get("Status"),
            "Current Status": cs.get("Status"),
        }
        for f in PCT_FIELDS:
            bv, cv = _num(bs.get(f)), _num(cs.get(f))
            row[f"{f} (baseline)"] = bv
            row[f"{f} (current)"] = cv
            row[f"{f} Δ"] = round(cv - bv, 2) if bv is not None and cv is not None else None
        for f in COUNT_FIELDS:
            row[f"{f} (baseline → current)"] = f"{bs.get(f)} → {cs.get(f)}"
        row["Home page changes"] = sum(1 for r in rows if r[2] == SOURCE_PAGE)
        row["Target page changes"] = sum(1 for r in rows if r[2] == TARGET_PAGE)
        overview.append(row)
        details += rows
    return overview, details


# ---------------------------------------------------------------------------
# Output
# ---------------------------------------------------------------------------

DETAIL_HEADERS = ["Component", "Locale", "Page", "Change", "Link / Text", "Path", "Baseline", "Current"]


def write_xlsx(path: Path, overview: list[dict], details: list[tuple], meta: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    headers: list[str] = []
    for r in overview:
        for k in r:
            if k not in headers:
                headers.append(k)
    ws.append(headers)
    style_header(ws)
    fills = {
        "Improved": PatternFill("solid", fgColor="C6EFCE"),
        "Regressed": PatternFill("solid", fgColor="FFC7CE"),
        "Error": PatternFill("solid", fgColor="FFC7CE"),
    }
    for i, r in enumerate(overview, 2):
        ws.append([r.get(h) for h in headers])
        fill = fills.get(r["Direction"], PatternFill("solid", fgColor="FFEB9C"))
        ws.cell(row=i, column=headers.index("Direction") + 1).fill = fill
    ws.freeze_panes = "C2"
    autofit(ws)

    wd = wb.create_sheet("Details")
    wd.append(DETAIL_HEADERS)
    style_header(wd)
    for d in details:
        wd.append(list(d))
    wd.freeze_panes = "A2"
    autofit(wd, max_width=80)

    wm = wb.create_sheet("Run Info")
    wm.append(["Field", "Value"])
    style_header(wm)
    for k, v in meta.items():
        wm.append([k, str(v)])
    autofit(wm, max_width=120)
    wb.save(path)


def write_html(path: Path, overview: list[dict], details: list[tuple], meta: dict):
    e = html.escape
    improved = sum(1 for r in overview if r["Direction"] == "Improved")
    regressed = sum(1 for r in overview if r["Direction"] in ("Regressed", "Error"))
    other = len(overview) - improved - regressed

    def pct_cell(r, f):
        bv, cv, dv = r.get(f"{f} (baseline)"), r.get(f"{f} (current)"), r.get(f"{f} Δ")
        if bv is None and cv is None:
            return "<td>—</td>"
        cls = "up" if (dv or 0) > 0 else "down" if (dv or 0) < 0 else ""
        delta = f"{dv:+.2f}" if dv else "0"
        return f'<td>{bv} → {cv} <span class="{cls}">({delta})</span></td>'

    over_rows = []
    for r in overview:
        cls = {"Improved": "up", "Regressed": "down", "Error": "down"}.get(r["Direction"], "same")
        over_rows.append(
            f"<tr><td>{e(r['Component'])}</td><td><a href=\"#{e(r['Component'])}-{e(r['Locale'])}\">"
            f"{e(r['Locale'])}</a></td><td><span class=\"tag {cls}\">{e(r['Direction'])}</span></td>"
            f"<td>{e(str(r.get('Baseline Status')))} → {e(str(r.get('Current Status')))}</td>"
            + "".join(pct_cell(r, f) for f in ("Overall Match %", "Structural Link Match %", "Text Match %"))
            + f"<td>{r.get('Home page changes', '—')}</td><td>{r.get('Target page changes', '—')}</td></tr>"
        )

    sections = []
    for r in overview:
        comp, loc = r["Component"], r["Locale"]
        rows = [d for d in details if d[0] == comp and d[1] == loc]
        blocks = []
        for page in (SOURCE_PAGE, TARGET_PAGE, "Comparison", "Run"):
            items = [d for d in rows if d[2] == page]
            if not items:
                continue
            lis = "".join(
                f"<tr><td class=\"chg\">{e(d[3])}</td><td>{e(str(d[4] or ''))}</td>"
                f"<td class=\"path\">{e(str(d[5] or ''))}</td>"
                f"<td class=\"path\">{e(str(d[6] or ''))}</td><td class=\"path\">{e(str(d[7] or ''))}</td></tr>"
                for d in items
            )
            blocks.append(
                f"<h4>{e(page)}</h4><table><tr><th>Change</th><th>Link / Text</th><th>Path</th>"
                f"<th>Baseline</th><th>Current</th></tr>{lis}</table>"
            )
        if not blocks:
            blocks.append('<p class="muted">Percentages changed, but no link or text difference was found on either page.</p>')
        sections.append(
            f"<section id=\"{e(comp)}-{e(loc)}\"><h3>{e(comp)} · {e(loc)} "
            f"<span class=\"tag {'up' if r['Direction']=='Improved' else 'down' if r['Direction'] in ('Regressed','Error') else 'same'}\">"
            f"{e(r['Direction'])}</span></h3>{''.join(blocks)}</section>"
        )

    doc = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Header Footer Difference Report</title>
<style>
:root{{--bg:#fff;--fg:#1d2433;--muted:#667085;--line:#e4e7ec;--card:#f8f9fb;--up:#127a3f;--upbg:#dcf5e5;--down:#b42318;--downbg:#fde4e2;--same:#8a5a00;--samebg:#fff2cc}}
@media (prefers-color-scheme:dark){{:root{{--bg:#14171c;--fg:#e6e8eb;--muted:#9aa3ad;--line:#2b3038;--card:#1b1f26;--up:#5fd08c;--upbg:#173527;--down:#ff8b80;--downbg:#3d1d1b;--same:#f0c060;--samebg:#3a301a}}}}
body{{margin:0;padding:24px 16px;background:var(--bg);color:var(--fg);font:14px/1.5 system-ui,Segoe UI,sans-serif}}
main{{max-width:1200px;margin:0 auto}} h1{{margin:0 0 4px}} .muted{{color:var(--muted)}}
.cards{{display:flex;gap:12px;flex-wrap:wrap;margin:16px 0}} .card{{background:var(--card);border:1px solid var(--line);border-radius:8px;padding:12px 16px;min-width:140px}}
.card b{{display:block;font-size:24px}} table{{border-collapse:collapse;width:100%;margin:8px 0 16px;display:block;overflow-x:auto}}
th,td{{border:1px solid var(--line);padding:6px 8px;text-align:left;vertical-align:top}} th{{background:var(--card)}}
.tag{{padding:2px 8px;border-radius:10px;font-size:12px;font-weight:600}} .tag.up{{background:var(--upbg);color:var(--up)}}
.tag.down{{background:var(--downbg);color:var(--down)}} .tag.same{{background:var(--samebg);color:var(--same)}}
span.up{{color:var(--up)}} span.down{{color:var(--down)}} .path{{font-family:Consolas,monospace;font-size:12px;word-break:break-all}}
section{{border-top:1px solid var(--line);padding-top:8px}} .chg{{white-space:nowrap}} a{{color:inherit}}
</style></head><body><main>
<h1>Header / Footer difference report</h1>
<p class="muted">Run {e(meta['Run folder'])} compared with baseline {e(meta['Baseline'])}.<br>
Home page = AEM (source) · Target page = Sitecore (press release page).</p>
<div class="cards"><div class="card"><b>{len(overview)}</b>locale/component rows changed</div>
<div class="card"><b class="up" style="color:var(--up)">{improved}</b>improved</div>
<div class="card"><b style="color:var(--down)">{regressed}</b>regressed / error</div>
<div class="card"><b>{other}</b>changed, overall % same</div></div>
<h2>Overview</h2>
<table><tr><th>Component</th><th>Locale</th><th>Direction</th><th>Status</th><th>Overall %</th>
<th>Link match %</th><th>Text match %</th><th>Home page changes</th><th>Target page changes</th></tr>
{''.join(over_rows)}</table>
<h2>Details</h2>{''.join(sections)}
</main></body></html>"""
    path.write_text(doc, encoding="utf-8")


def run(run_dir: Path, baseline_dir: Path = BASELINE_DIR) -> dict:
    """Diff run_dir against baseline_dir. Returns a small summary dict."""
    overview, details = [], []
    for component, fname in COMPONENTS:
        base_path, cur_path = baseline_dir / fname, run_dir / fname
        if not base_path.exists() or not cur_path.exists():
            print(f"  {component}: report missing (baseline={base_path.exists()} current={cur_path.exists()})")
            continue
        o, d = diff_component(component, load_report(base_path), load_report(cur_path))
        overview += o
        details += d

    summary = {
        "changed": len(overview),
        "improved": sum(1 for r in overview if r["Direction"] == "Improved"),
        "regressed": sum(1 for r in overview if r["Direction"] in ("Regressed", "Error")),
        "report": None,
    }
    if not overview:
        print("No change vs baseline — difference report not generated.")
        return summary

    meta = {
        "Run folder": str(run_dir),
        "Baseline": str(baseline_dir),
        "Generated": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "Trigger": "Any change in a match % (any size), status, or links/texts on either page",
    }
    xlsx, html_path = run_dir / "DifferenceReport.xlsx", run_dir / "DifferenceReport.html"
    write_xlsx(xlsx, overview, details, meta)
    write_html(html_path, overview, details, meta)
    summary["report"] = str(html_path)
    print(
        f"CHANGES vs baseline: {summary['changed']} locale/component rows "
        f"(improved={summary['improved']}, regressed={summary['regressed']})"
    )
    for r in overview:
        print(
            f"  {r['Component']:<6} {r['Locale']:<6} {r['Direction']:<26} "
            f"overall {r.get('Overall Match % (baseline)')} -> {r.get('Overall Match % (current)')}"
        )
    print(f"Difference report: {html_path}")
    return summary


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("run_dir")
    ap.add_argument("--baseline", default=str(BASELINE_DIR))
    args = ap.parse_args()
    run(Path(args.run_dir), Path(args.baseline))


if __name__ == "__main__":
    main()
